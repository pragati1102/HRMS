import datetime
import calendar
import csv
import ipaddress
import mimetypes
from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from django.contrib.auth.views import PasswordChangeView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import TemplateView
from django.views import View
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse, reverse_lazy
from django.http import FileResponse, Http404, HttpResponse
from django.core.paginator import Paginator
from django.db.models import Q
from django.contrib import messages
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .forms import EmployeeProfileForm
from .models import User, Attendance, AttendanceCorrectionRequest, CompanyHoliday, WorkLog, LeaveApplication, Salary, Project, ProjectAssignment, EmployeeDocument
from .permissions import AdminRoleRequiredMixin, EmployeeRoleRequiredMixin

def get_leave_days_in_range(employee, start_limit, end_limit):
    """
    Calculates the total days of approved leave in the specified range.
    """
    leaves = LeaveApplication.objects.filter(
        employee=employee,
        status='approved',
        start_date__lte=end_limit,
        end_date__gte=start_limit
    )
    holidays = set(CompanyHoliday.objects.filter(date__range=(start_limit, end_limit)).values_list('date', flat=True))
    total_days = Decimal('0')
    for leave in leaves:
        overlap_start = max(leave.start_date, start_limit)
        overlap_end = min(leave.end_date, end_limit)
        total_days += get_leave_units(overlap_start, overlap_end, leave.half_day_session, holidays)
    return total_days


def get_leave_units(start_date, end_date, half_day_session='full', holidays=None):
    holidays = holidays or set()
    if half_day_session != 'full':
        return Decimal('0.5') if start_date == end_date and start_date.weekday() < 5 and start_date not in holidays else Decimal('0')

    units = Decimal('0')
    current_date = start_date
    while current_date <= end_date:
        if current_date.weekday() < 5 and current_date not in holidays:
            units += Decimal('1')
        current_date += datetime.timedelta(days=1)
    return units


def get_leave_quota_summary(employee, year):
    year_start = datetime.date(year, 1, 1)
    year_end = datetime.date(year, 12, 31)
    holidays = set(CompanyHoliday.objects.filter(date__range=(year_start, year_end)).values_list('date', flat=True))
    quota_fields = {
        'Casual': employee.casual_leave_quota,
        'Sick': employee.sick_leave_quota,
        'Paid': employee.paid_leave_quota,
    }
    summaries = []
    for leave_type, quota in quota_fields.items():
        used = Decimal('0')
        pending = Decimal('0')
        applications = LeaveApplication.objects.filter(
            employee=employee,
            leave_type=leave_type,
            status__in=['approved', 'pending'],
            start_date__lte=year_end,
            end_date__gte=year_start,
        )
        for application in applications:
            overlap_start = max(application.start_date, year_start)
            overlap_end = min(application.end_date, year_end)
            units = get_leave_units(overlap_start, overlap_end, application.half_day_session, holidays)
            if application.status == 'approved':
                used += units
            else:
                pending += units
        available = max(Decimal('0'), Decimal(quota) - used - pending) if quota is not None else None
        summaries.append({
            'leave_type': leave_type,
            'quota': quota,
            'used': used,
            'pending': pending,
            'available': available,
        })
    return summaries


def get_next_worklog_slot(employee, date_value):
    attendance = Attendance.objects.filter(employee=employee, date=date_value).first()
    if not attendance or not attendance.check_in:
        return None, None

    candidate = datetime.datetime.combine(date_value, attendance.check_in)
    shift_end = (
        datetime.datetime.combine(date_value, attendance.check_out)
        if attendance.check_out else None
    )
    logs = [
        (
            datetime.datetime.combine(date_value, log.start_time),
            datetime.datetime.combine(date_value, log.end_time),
        )
        for log in WorkLog.objects.filter(employee=employee, date=date_value)
    ]

    while shift_end is None or candidate < shift_end:
        slot_end = candidate + datetime.timedelta(hours=1)
        if shift_end:
            slot_end = min(slot_end, shift_end)
        if not any(log_start < slot_end and log_end > candidate for log_start, log_end in logs):
            return candidate.time(), slot_end.time()
        candidate += datetime.timedelta(hours=1)
    return None, None


def build_worklog_workbook(employee):
    workbook = Workbook()
    logs_by_month = defaultdict(list)
    logs = list(WorkLog.objects.filter(employee=employee).order_by('date', 'start_time'))
    for log in logs:
        logs_by_month[(log.date.year, log.date.month)].append(log)

    today = timezone.localdate()
    logs_by_month.setdefault((today.year, today.month), [])
    attendance_by_date = {
        item.date: item
        for item in Attendance.objects.filter(
            employee=employee,
            date__in={log.date for log in logs},
        )
    }
    header_fill = PatternFill('solid', fgColor='172554')
    stripe_fill = PatternFill('solid', fgColor='F1F5F9')
    total_fill = PatternFill('solid', fgColor='DBEAFE')
    thin_border = Border(bottom=Side(style='thin', color='CBD5E1'))
    header_font = Font(bold=True, color='FFFFFF')

    for sheet_index, ((year, month), month_logs) in enumerate(sorted(logs_by_month.items())):
        sheet = workbook.active if sheet_index == 0 else workbook.create_sheet()
        month_date = datetime.date(year, month, 1)
        sheet.title = month_date.strftime('%b %Y')
        sheet['A1'] = 'Employee Name'
        sheet['B1'] = employee.get_full_name() or employee.username
        sheet['D1'] = 'Employee ID'
        sheet['E1'] = employee.pk
        sheet['A2'] = 'Department'
        sheet['B2'] = employee.department or '-'
        sheet['D2'] = 'Month / Year'
        sheet['E2'] = month_date.strftime('%B %Y')
        for cell in ('A1', 'D1', 'A2', 'D2'):
            sheet[cell].font = Font(bold=True, color='172554')

        headers_row = 4
        headers = ['Date', 'Login Time', 'Time Slot', 'Category', 'Task Description', 'Total Hours', 'Logout Time']
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row=headers_row, column=column, value=value)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border

        rows_by_date = defaultdict(list)
        for log in month_logs:
            rows_by_date[log.date].append(log)

        row_number = headers_row + 1
        daily_total_rows = []
        stripe_index = 0
        for date_value, date_logs in sorted(rows_by_date.items()):
            detail_start = row_number
            attendance = attendance_by_date.get(date_value)
            for log in date_logs:
                slot = f'{log.start_time.strftime("%I:%M %p")} - {log.end_time.strftime("%I:%M %p")}'
                values = [
                    date_value,
                    attendance.check_in if attendance and attendance.check_in else None,
                    slot,
                    log.category,
                    log.description,
                    float(log.hours_worked or 0),
                    attendance.check_out if attendance and attendance.check_out else None,
                ]
                for column, value in enumerate(values, start=1):
                    cell = sheet.cell(row=row_number, column=column, value=value)
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical='top', wrap_text=column == 5)
                    if stripe_index % 2:
                        cell.fill = stripe_fill
                sheet.cell(row=row_number, column=1).number_format = 'mmm d, yyyy'
                for column in (2, 7):
                    sheet.cell(row=row_number, column=column).number_format = 'h:mm AM/PM'
                sheet.cell(row=row_number, column=6).number_format = '0.00'
                row_number += 1
                stripe_index += 1

            daily_total_row = row_number
            sheet.cell(row=daily_total_row, column=5, value='Daily Total')
            sheet.cell(row=daily_total_row, column=6, value=f'=SUM(F{detail_start}:F{daily_total_row - 1})')
            for cell in sheet[daily_total_row]:
                cell.fill = total_fill
                cell.font = Font(bold=True, color='172554')
                cell.border = thin_border
            sheet.cell(row=daily_total_row, column=6).number_format = '0.00'
            daily_total_rows.append(daily_total_row)
            row_number += 1

        month_total_row = row_number + 1
        sheet.cell(row=month_total_row, column=5, value='Month Total')
        total_formula = '=SUM(' + ','.join(f'F{row}' for row in daily_total_rows) + ')' if daily_total_rows else '=0'
        sheet.cell(row=month_total_row, column=6, value=total_formula)
        for cell in sheet[month_total_row]:
            cell.fill = total_fill
            cell.font = Font(bold=True, color='172554')
            cell.border = thin_border
        sheet.cell(row=month_total_row, column=6).number_format = '0.00'
        sheet.freeze_panes = 'A5'
        sheet.auto_filter.ref = f'A{headers_row}:G{max(headers_row, row_number - 1)}'

        for column in range(1, len(headers) + 1):
            values = [len(str(sheet.cell(row, column).value or '')) for row in range(1, month_total_row + 1)]
            width = max(values + [len(headers[column - 1])]) + 2
            sheet.column_dimensions[get_column_letter(column)].width = min(width, 50)
        sheet.column_dimensions['E'].width = max(sheet.column_dimensions['E'].width, 36)
        sheet.row_dimensions[headers_row].height = 24

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


class EmployeeDashboardView(EmployeeRoleRequiredMixin, TemplateView):
    """
    Dashboard for the logged-in employee. Shows module quick links and
    a feed of recent work log activity.
    """
    template_name = 'core/employee_dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user

        pending_applications_count = LeaveApplication.objects.filter(employee=user, status='pending').count()
        recent_activity = WorkLog.objects.filter(employee=user).order_by('-date', '-start_time')[:5]

        context.update({
            'pending_applications_count': pending_applications_count,
            'recent_activity': recent_activity,
        })
        return context


class EmployeeCheckInView(EmployeeRoleRequiredMixin, View):
    """
    Marks today's attendance as present and logs check-in time.
    """
    def post(self, request):
        today = timezone.localdate()
        location_mode = request.POST.get('location_mode', '').strip()
        if location_mode not in ('office', 'remote'):
            messages.error(request, 'Choose Office or Remote before checking in.')
            return redirect('employee_attendance')

        remote_ip = request.META.get('REMOTE_ADDR', '').strip()
        try:
            remote_ip = str(ipaddress.ip_address(remote_ip))
        except ValueError:
            remote_ip = None

        att = Attendance.objects.filter(employee=request.user, date=today).first()
        if att and att.check_in:
            messages.warning(request, "You have already checked in today.")
            return redirect(request.META.get('HTTP_REFERER', 'employee_attendance'))
            
        att, created = Attendance.objects.get_or_create(
            employee=request.user, 
            date=today,
            defaults={
                'status': 'present',
                'check_in': timezone.localtime().time(),
                'location_mode': location_mode,
                'check_in_ip': remote_ip,
            }
        )
        if not created:
            if att.status in ['absent', 'leave', 'no_record']:
                att.status = 'present'
            if not att.check_in:
                att.check_in = timezone.localtime().time()
            att.location_mode = location_mode
            att.check_in_ip = remote_ip
            att.save()
            messages.info(request, "Today's check-in has been logged.")
        else:
            messages.success(request, "Checked in successfully!")
        return redirect(request.META.get('HTTP_REFERER', 'employee_attendance'))


class EmployeeCheckOutView(EmployeeRoleRequiredMixin, View):
    """
    Logs check-out time for today's attendance.
    """
    def post(self, request):
        today = timezone.localdate()
        att = Attendance.objects.filter(employee=request.user, date=today).first()
        if att:
            if att.check_out:
                messages.warning(request, "You have already checked out today.")
            else:
                now = timezone.now()
                if att.break_started_at:
                    att.break_seconds += max(0, int((now - att.break_started_at).total_seconds()))
                    att.break_started_at = None
                att.check_out = timezone.localtime().time()
                att.save(update_fields=['check_out', 'break_seconds', 'break_started_at'])
                messages.success(request, "Checked out successfully!")
        else:
            messages.error(request, "No attendance record found for today. Please check-in first.")
        return redirect(request.META.get('HTTP_REFERER', 'employee_attendance'))


class EmployeeAttendanceBreakView(EmployeeRoleRequiredMixin, View):
    def post(self, request):
        today = timezone.localdate()
        attendance = Attendance.objects.filter(employee=request.user, date=today).first()
        if not attendance or not attendance.check_in or attendance.check_out:
            messages.error(request, 'A break can only be recorded during an active work session.')
            return redirect('employee_attendance')

        if attendance.break_started_at:
            attendance.break_seconds += max(0, int((timezone.now() - attendance.break_started_at).total_seconds()))
            attendance.break_started_at = None
            messages.success(request, 'Break ended. Your work timer has resumed.')
        else:
            attendance.break_started_at = timezone.now()
            messages.success(request, 'Break started.')
        attendance.save(update_fields=['break_seconds', 'break_started_at'])
        return redirect('employee_attendance')


class EmployeeAttendanceCorrectionView(EmployeeRoleRequiredMixin, View):
    def post(self, request):
        date_value = request.POST.get('date', '').strip()
        reason = request.POST.get('reason', '').strip()
        check_in_value = request.POST.get('requested_check_in', '').strip()
        check_out_value = request.POST.get('requested_check_out', '').strip()

        try:
            correction_date = datetime.date.fromisoformat(date_value)
            requested_check_in = datetime.time.fromisoformat(check_in_value) if check_in_value else None
            requested_check_out = datetime.time.fromisoformat(check_out_value) if check_out_value else None
        except ValueError:
            messages.error(request, 'Enter a valid date and time for the correction request.')
            return redirect('employee_attendance')

        if not reason or (not requested_check_in and not requested_check_out):
            messages.error(request, 'Add a reason and at least one corrected time.')
            return redirect('employee_attendance')
        if correction_date > timezone.localdate():
            messages.error(request, 'Attendance corrections cannot be requested for a future date.')
            return redirect('employee_attendance')

        AttendanceCorrectionRequest.objects.create(
            employee=request.user,
            date=correction_date,
            requested_check_in=requested_check_in,
            requested_check_out=requested_check_out,
            reason=reason,
        )
        messages.success(request, 'Attendance correction request sent to HR.')
        return redirect('employee_attendance')


class EmployeeAttendanceExportView(EmployeeRoleRequiredMixin, View):
    def get(self, request):
        today = timezone.localdate()
        month = request.GET.get('month', str(today.month))
        year = request.GET.get('year', str(today.year))
        try:
            month_number = int(month)
            year_number = int(year)
            if month_number not in range(1, 13):
                raise ValueError
        except ValueError:
            return HttpResponse('Invalid month or year.', status=400, content_type='text/plain')

        records = Attendance.objects.filter(
            employee=request.user,
            date__month=month_number,
            date__year=year_number,
        ).order_by('date')
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="attendance-{year_number}-{month_number:02d}.csv"'
        writer = csv.writer(response)
        writer.writerow(['Date', 'Status', 'Check In', 'Check Out', 'Working Hours', 'Break Duration', 'Mode', 'IP', 'Remarks'])
        for record in records:
            working_seconds = get_attendance_work_seconds(record)
            writer.writerow([
                record.date.isoformat(),
                record.get_status_display(),
                record.check_in or '',
                record.check_out or '',
                format_duration(working_seconds),
                format_duration(record.break_seconds),
                record.location_mode,
                record.check_in_ip or '',
                record.remarks,
            ])
        return response


def get_attendance_work_seconds(attendance, now=None):
    if not attendance.check_in:
        return 0
    start = datetime.datetime.combine(attendance.date, attendance.check_in)
    if attendance.check_out:
        end = datetime.datetime.combine(attendance.date, attendance.check_out)
        if end < start:
            end += datetime.timedelta(days=1)
        gross_seconds = int((end - start).total_seconds())
    elif attendance.date == timezone.localdate():
        current_time = timezone.localtime(now or timezone.now())
        end = datetime.datetime.combine(attendance.date, current_time.time().replace(tzinfo=None))
        if end < start:
            end += datetime.timedelta(days=1)
        gross_seconds = int((end - start).total_seconds())
    else:
        return 0

    break_seconds = attendance.break_seconds
    if attendance.break_started_at:
        break_seconds += max(0, int(((now or timezone.now()) - attendance.break_started_at).total_seconds()))
    return max(0, gross_seconds - break_seconds)


def format_duration(seconds):
    hours, remainder = divmod(max(0, seconds), 3600)
    minutes = remainder // 60
    return f'{hours}h {minutes:02d}m' if hours else f'{minutes}m'


class EmployeeAttendanceView(EmployeeRoleRequiredMixin, TemplateView):
    """
    Shows today's attendance marking options, monthly summary stats,
    and history table filterable by month/year.
    """
    template_name = 'core/attendance.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        today = timezone.localdate()
        
        # Month/Year Filter (default to current month/year)
        month_str = self.request.GET.get('month')
        year_str = self.request.GET.get('year')
        
        selected_month = int(month_str) if month_str and month_str.isdigit() and 1 <= int(month_str) <= 12 else today.month
        selected_year = int(year_str) if year_str and year_str.isdigit() and 2000 <= int(year_str) <= today.year + 3 else today.year
        
        # Filtered history records
        attendances = Attendance.objects.filter(
            employee=user, 
            date__month=selected_month, 
            date__year=selected_year
        ).order_by('-date')
        
        # Monthly statistics based on filtered month
        total_present = attendances.filter(status='present').count()
        total_absent = attendances.filter(status='absent').count()
        total_half_days = attendances.filter(status='half_day').count()
        expected_start = datetime.time(9, 15)
        expected_end = datetime.time(17, 0)
        total_work_seconds = 0
        total_overtime_seconds = 0
        late_arrivals = 0
        early_checkouts = 0
        days_with_hours = 0
        attendance_rows = []
        correction_dates = set(
            request.date for request in user.attendance_corrections.filter(status='pending')
        )
        now = timezone.now()
        for attendance in attendances:
            work_seconds = get_attendance_work_seconds(attendance, now=now)
            total_work_seconds += work_seconds
            if attendance.check_in and attendance.check_out:
                days_with_hours += 1
                total_overtime_seconds += max(0, work_seconds - 8 * 3600)
            if attendance.check_in and attendance.check_in > expected_start:
                late_arrivals += 1
            if attendance.check_out and attendance.check_out < expected_end:
                early_checkouts += 1
            attendance.work_seconds = work_seconds
            attendance.formatted_work_time = format_duration(work_seconds)
            attendance.formatted_break_time = format_duration(attendance.break_seconds)
            attendance.status_badge_class = {
                'present': 'bg-success-subtle text-success',
                'half_day': 'bg-warning-subtle text-warning',
                'absent': 'bg-danger-subtle text-danger',
                'leave': 'bg-info-subtle text-info',
            }.get(attendance.status, 'bg-secondary-subtle text-secondary')
            attendance_rows.append(attendance)
        
        # Generate filter selections
        months = [(i, calendar.month_name[i]) for i in range(1, 13)]
        years = range(today.year - 3, today.year + 3)
        
        # Today's attendance record
        today_attendance = Attendance.objects.filter(employee=user, date=today).first()
        current_work_seconds = get_attendance_work_seconds(today_attendance, now=timezone.now()) if today_attendance else 0
        
        context.update({
            'attendances': attendances,
            'selected_month': selected_month,
            'selected_year': selected_year,
            'total_present': total_present,
            'total_absent': total_absent,
            'total_half_days': total_half_days,
            'total_work_time': format_duration(total_work_seconds),
            'average_work_time': format_duration(total_work_seconds // days_with_hours) if days_with_hours else '0m',
            'late_arrivals': late_arrivals,
            'early_checkouts': early_checkouts,
            'overtime_time': format_duration(total_overtime_seconds),
            'attendance_rows': attendance_rows,
            'pending_correction_dates': correction_dates,
            'current_time': timezone.localtime(),
            'months': months,
            'years': years,
            'today_attendance': today_attendance,
            'today': today,
            'current_work_seconds': current_work_seconds,
            'current_work_time': format_duration(current_work_seconds),
        })
        return context


class EmployeeWorkLogsView(EmployeeRoleRequiredMixin, TemplateView):
    """
    Manages daily work logging: form to add slots, today's log listing with edit/delete,
    and history list filterable by date range grouped by day.
    """
    template_name = 'core/work_logs.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        today = timezone.localdate()
        suggested_start_time, suggested_end_time = get_next_worklog_slot(user, today)
        
        # Today's logs
        today_work_logs = WorkLog.objects.filter(employee=user, date=today).order_by('start_time')
        total_hours_today = sum(log.hours_worked for log in today_work_logs) if today_work_logs else 0
        
        # Date range filters
        start_date_str = self.request.GET.get('start_date', '')
        end_date_str = self.request.GET.get('end_date', '')
        
        history_logs = WorkLog.objects.filter(employee=user).order_by('-date', '-start_time')
        
        if start_date_str:
            try:
                start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date()
                history_logs = history_logs.filter(date__gte=start_date)
            except ValueError:
                pass
        if end_date_str:
            try:
                end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date()
                history_logs = history_logs.filter(date__lte=end_date)
            except ValueError:
                pass
                
        # Group history in Python
        logs_by_day = {}
        for log in history_logs:
            logs_by_day.setdefault(log.date, []).append(log)
            
        grouped_history = []
        for date in sorted(logs_by_day.keys(), reverse=True):
            day_logs = logs_by_day[date]
            total_hours = sum(l.hours_worked for l in day_logs)
            grouped_history.append({
                'date': date,
                'logs': day_logs,
                'total_hours': total_hours
            })
            
        context.update({
            'today_work_logs': today_work_logs,
            'total_hours_today': total_hours_today,
            'grouped_history': grouped_history,
            'start_date': start_date_str,
            'end_date': end_date_str,
            'today': today,
            'categories': ['Meeting', 'Development', 'Bug Fix', 'Documentation', 'Testing', 'Other'],
            'suggested_start_time': suggested_start_time,
            'suggested_end_time': suggested_end_time,
            'today_attendance': Attendance.objects.filter(employee=user, date=today).first(),
        })
        return context


class EmployeeWorkLogExportView(EmployeeRoleRequiredMixin, View):
    def get(self, request):
        output = build_worklog_workbook(request.user)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{request.user.username}_task_management.xlsx"'
        return response


class AdminWorkLogExportView(AdminRoleRequiredMixin, View):
    def get(self, request, employee_id):
        employee = get_object_or_404(User, pk=employee_id, role='employee')
        output = build_worklog_workbook(employee)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{employee.username}_task_management.xlsx"'
        return response


class EmployeeAddWorkLogView(EmployeeRoleRequiredMixin, View):
    """
    Adds a new work log entry. Validates date, range, and slot overlaps.
    """
    def post(self, request):
        date_str = request.POST.get('date')
        start_time_str = request.POST.get('start_time')
        end_time_str = request.POST.get('end_time')
        description = request.POST.get('description', '').strip()
        category = request.POST.get('category', 'Other')
        
        if not start_time_str or not end_time_str or not description:
            messages.error(request, "Start time, end time, and description are required.")
            return redirect('employee_work_logs')
            
        try:
            if date_str:
                log_date = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
            else:
                log_date = timezone.localdate()
                
            start_time = datetime.datetime.strptime(start_time_str, '%H:%M').time()
            end_time = datetime.datetime.strptime(end_time_str, '%H:%M').time()
            valid_categories = dict(WorkLog.CATEGORY_CHOICES)
            attendance = Attendance.objects.filter(employee=request.user, date=log_date).first()

            if category not in valid_categories:
                messages.error(request, 'Choose a valid task category.')
                return redirect('employee_work_logs')

            if not attendance or not attendance.check_in:
                messages.error(request, 'Check in before logging tasks for this date.')
                return redirect('employee_work_logs')
            
            if start_time >= end_time:
                messages.error(request, "End time must be after start time.")
                return redirect('employee_work_logs')

            if start_time < attendance.check_in or (attendance.check_out and end_time > attendance.check_out):
                messages.error(request, 'Task times must be within your recorded shift.')
                return redirect('employee_work_logs')
                
            # Overlap check
            overlap = WorkLog.objects.filter(
                employee=request.user,
                date=log_date,
                start_time__lt=end_time,
                end_time__gt=start_time
            ).exists()
            
            if overlap:
                messages.error(request, "Error: This slot overlaps with another work log entry for the same day.")
                return redirect('employee_work_logs')
                
            WorkLog.objects.create(
                employee=request.user,
                date=log_date,
                start_time=start_time,
                end_time=end_time,
                description=description,
                category=category
            )
            messages.success(request, "Work log entry added successfully!")
        except Exception as e:
            messages.error(request, f"Error saving work log: {str(e)}")
            
        return redirect('employee_work_logs')


class EmployeeDeleteWorkLogView(EmployeeRoleRequiredMixin, View):
    """
    Deletes a work log entry only if it belongs to the current day.
    """
    def post(self, request, pk):
        log = get_object_or_404(WorkLog, pk=pk, employee=request.user)
        today = timezone.localdate()
        if log.date != today:
            messages.error(request, "You can only delete work log entries created today.")
        else:
            log.delete()
            messages.success(request, "Work log entry deleted successfully.")
        return redirect('employee_work_logs')


class EmployeeEditWorkLogView(EmployeeRoleRequiredMixin, View):
    """
    Renders edit page and handles edit submission. Restricted to current day logs.
    """
    template_name = 'core/edit_work_log.html'

    def get(self, request, pk):
        log = get_object_or_404(WorkLog, pk=pk, employee=request.user)
        today = timezone.localdate()
        if log.date != today:
            messages.error(request, "You can only edit work log entries created today.")
            return redirect('employee_work_logs')
        
        categories = ['Meeting', 'Development', 'Bug Fix', 'Documentation', 'Testing', 'Other']
        return render(request, self.template_name, {'log': log, 'categories': categories})

    def post(self, request, pk):
        log = get_object_or_404(WorkLog, pk=pk, employee=request.user)
        today = timezone.localdate()
        if log.date != today:
            messages.error(request, "You can only edit work log entries created today.")
            return redirect('employee_work_logs')
            
        start_time_str = request.POST.get('start_time')
        end_time_str = request.POST.get('end_time')
        description = request.POST.get('description', '').strip()
        category = request.POST.get('category', 'Other')
        
        if not start_time_str or not end_time_str or not description:
            messages.error(request, "All fields are required.")
            return redirect('employee_edit_worklog', pk=pk)
            
        try:
            start_time = datetime.datetime.strptime(start_time_str, '%H:%M').time()
            end_time = datetime.datetime.strptime(end_time_str, '%H:%M').time()
            
            if start_time >= end_time:
                messages.error(request, "End time must be after start time.")
                return redirect('employee_edit_worklog', pk=pk)
                
            overlap = WorkLog.objects.filter(
                employee=request.user,
                date=log.date,
                start_time__lt=end_time,
                end_time__gt=start_time
            ).exclude(id=log.id).exists()
            
            if overlap:
                messages.error(request, "Error: This slot overlaps with another work log entry.")
                return redirect('employee_edit_worklog', pk=pk)
                
            log.start_time = start_time
            log.end_time = end_time
            log.description = description
            log.category = category
            log.save()
            
            messages.success(request, "Work log updated successfully!")
            return redirect('employee_work_logs')
        except Exception as e:
            messages.error(request, f"Error saving work log: {str(e)}")
            return redirect('employee_edit_worklog', pk=pk)


class EmployeeLeavesView(EmployeeRoleRequiredMixin, TemplateView):
    """
    Renders the leave application form and list of personal applications.
    """
    template_name = 'core/leaves.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        today = timezone.localdate()
        
        # Statistics (Annual total leaves, pending applications)
        first_day_of_year = datetime.date(today.year, 1, 1)
        last_day_of_year = datetime.date(today.year, 12, 31)
        leaves_taken_year = get_leave_days_in_range(user, first_day_of_year, last_day_of_year)
        pending_applications_count = LeaveApplication.objects.filter(employee=user, status='pending').count()
        
        leave_type_filter = self.request.GET.get('status', 'all').lower()
        if leave_type_filter not in {'all', 'pending', 'approved', 'rejected'}:
            leave_type_filter = 'all'
        search_query = self.request.GET.get('q', '').strip()[:100]
        applications = LeaveApplication.objects.filter(employee=user).select_related('approved_by', 'handover_contact')
        if leave_type_filter != 'all':
            applications = applications.filter(status=leave_type_filter)
        if search_query:
            applications = applications.filter(Q(reason__icontains=search_query) | Q(leave_type__icontains=search_query))
        paginator = Paginator(applications.order_by('-applied_on'), 8)
        leave_applications_page = paginator.get_page(self.request.GET.get('page'))
        leave_quota_summary = get_leave_quota_summary(user, today.year)
        leave_balance_by_type = {
            item['leave_type']: item['available'] for item in leave_quota_summary
        }
        used_leave_by_type = {
            item['leave_type']: item['used'] for item in leave_quota_summary
        }
        leave_filter_counts = {
            status: LeaveApplication.objects.filter(employee=user, status=status).count()
            for status in ['pending', 'approved', 'rejected']
        }
        
        context.update({
            'leaves_taken_year': leaves_taken_year,
            'pending_applications_count': pending_applications_count,
            'leave_applications': leave_applications_page,
            'leave_quota_summary': leave_quota_summary,
            'leave_balance_by_type': leave_balance_by_type,
            'used_leave_by_type': used_leave_by_type,
            'leave_filter': leave_type_filter,
            'leave_filter_counts': leave_filter_counts,
            'search_query': search_query,
            'colleagues': User.objects.filter(role='employee', is_active=True).exclude(pk=user.pk).order_by('first_name', 'last_name'),
            'holiday_dates': [date.isoformat() for date in CompanyHoliday.objects.filter(date__year=today.year).values_list('date', flat=True)],
            'leave_types': [choice[0] for choice in LeaveApplication.LEAVE_TYPE_CHOICES],
            'today': today,
        })
        return context


class EmployeeApplyLeaveView(EmployeeRoleRequiredMixin, View):
    """
    Handles submission of new leave applications. Validates start and end dates.
    """
    def post(self, request):
        start_date_str = request.POST.get('start_date')
        end_date_str = request.POST.get('end_date')
        reason = request.POST.get('reason', '').strip()
        leave_type = request.POST.get('leave_type', 'Casual')
        half_day_session = request.POST.get('half_day_session', 'full')
        handover_id = request.POST.get('handover_contact', '').strip()
        supporting_document = request.FILES.get('supporting_document')
        
        if not start_date_str or not end_date_str or not reason:
            messages.error(request, "Start date, end date, and reason are required.")
            return redirect('employee_leaves')
            
        try:
            start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date()
            end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date()
            today = timezone.localdate()
            valid_leave_types = dict(LeaveApplication.LEAVE_TYPE_CHOICES)
            valid_sessions = dict(LeaveApplication.HALF_DAY_CHOICES)
            
            if start_date < today:
                messages.error(request, "Leave start date cannot be in the past.")
                return redirect('employee_leaves')
                
            if start_date > end_date:
                messages.error(request, "End date cannot be before start date.")
                return redirect('employee_leaves')

            if start_date.year != end_date.year:
                messages.error(request, "A leave request cannot span calendar years. Submit separate requests for each year.")
                return redirect('employee_leaves')

            if leave_type not in valid_leave_types or half_day_session not in valid_sessions:
                messages.error(request, "Select a valid leave type and leave session.")
                return redirect('employee_leaves')

            if half_day_session != 'full' and start_date != end_date:
                messages.error(request, "Half-day requests must cover a single date.")
                return redirect('employee_leaves')

            if leave_type in {'Sick', 'Maternity', 'Paternity'} and not supporting_document:
                messages.error(request, "A supporting document is required for this leave type.")
                return redirect('employee_leaves')

            if supporting_document:
                allowed_extensions = {'.pdf', '.jpg', '.jpeg', '.png'}
                if Path(supporting_document.name).suffix.lower() not in allowed_extensions or supporting_document.size > 10 * 1024 * 1024:
                    messages.error(request, "Upload a PDF or image no larger than 10 MB.")
                    return redirect('employee_leaves')

            holiday_dates = set(CompanyHoliday.objects.filter(date__range=(start_date, end_date)).values_list('date', flat=True))
            requested_units = get_leave_units(start_date, end_date, half_day_session, holiday_dates)
            if requested_units <= 0:
                messages.error(request, "The selected dates contain no working days.")
                return redirect('employee_leaves')

            quota_field = {'Casual': 'casual_leave_quota', 'Sick': 'sick_leave_quota', 'Paid': 'paid_leave_quota'}.get(leave_type)
            quota = getattr(request.user, quota_field) if quota_field else None
            if quota is not None:
                existing_units = Decimal('0')
                existing_applications = LeaveApplication.objects.filter(
                    employee=request.user,
                    leave_type=leave_type,
                    status__in=['approved', 'pending'],
                    start_date__year=start_date.year,
                )
                for application in existing_applications:
                    existing_holidays = set(CompanyHoliday.objects.filter(date__range=(application.start_date, application.end_date)).values_list('date', flat=True))
                    existing_units += get_leave_units(application.start_date, application.end_date, application.half_day_session, existing_holidays)
                available_units = max(Decimal('0'), Decimal(quota) - existing_units)
                if requested_units > available_units:
                    messages.error(request, f"This request exceeds your available {leave_type.lower()} leave balance ({available_units} days remaining).")
                    return redirect('employee_leaves')

            handover_contact = None
            if handover_id:
                handover_contact = get_object_or_404(User, pk=handover_id, role='employee', is_active=True)
                if handover_contact.pk == request.user.pk:
                    messages.error(request, "Choose a different colleague for handover.")
                    return redirect('employee_leaves')
                
            LeaveApplication.objects.create(
                employee=request.user,
                start_date=start_date,
                end_date=end_date,
                reason=reason,
                leave_type=leave_type,
                status='pending',
                half_day_session=half_day_session,
                supporting_document=supporting_document,
                handover_contact=handover_contact,
            )
            messages.success(request, "Leave application submitted successfully!")
        except ValueError:
            messages.error(request, "Enter valid start and end dates.")
            
        return redirect('employee_leaves')


class EmployeeWithdrawLeaveView(EmployeeRoleRequiredMixin, View):
    def post(self, request, pk):
        application = get_object_or_404(LeaveApplication, pk=pk, employee=request.user)
        if application.status != 'pending':
            messages.error(request, "Only pending leave applications can be withdrawn.")
        else:
            application.status = 'withdrawn'
            application.save(update_fields=['status'])
            messages.success(request, "Leave application withdrawn.")
        return redirect('employee_leaves')


class EmployeeLeaveDocumentView(EmployeeRoleRequiredMixin, View):
    def get(self, request, pk):
        application = get_object_or_404(LeaveApplication, pk=pk, employee=request.user)
        if not application.supporting_document:
            raise Http404
        return FileResponse(
            application.supporting_document.open('rb'),
            as_attachment=True,
            filename=Path(application.supporting_document.name).name,
            content_type='application/octet-stream',
        )


class EmployeeProfileView(EmployeeRoleRequiredMixin, View):
    """
    Renders personal profile details and allows updating gender and contact details.
    """
    template_name = 'core/profile.html'

    def get(self, request):
        return self.render_profile(request)

    def render_profile(self, request, form=None):
        user = request.user
        today = timezone.localdate()
        first_day_of_year = datetime.date(today.year, 1, 1)
        last_day_of_year = datetime.date(today.year, 12, 31)
        leave_taken = get_leave_days_in_range(user, first_day_of_year, last_day_of_year)
        active_projects = user.projects.filter(status='active').order_by('name')
        latest_leave = user.leave_applications.order_by('-applied_on').first()
        latest_work_log = user.work_logs.order_by('-date', '-start_time').first()
        years_of_service = 0
        if user.joining_date:
            years_of_service = today.year - user.joining_date.year - (
                (today.month, today.day) < (user.joining_date.month, user.joining_date.day)
            )

        context = {
            'profile_form': form or EmployeeProfileForm(instance=user),
            'profile_edit_mode': form is not None,
            'active_projects': active_projects,
            'latest_leave': latest_leave,
            'latest_work_log': latest_work_log,
            'years_of_service': years_of_service,
            'leave_balance': max(0, user.leave_entitlement_days - leave_taken) if user.leave_entitlement_days is not None else None,
            'masked_bank_account': self.mask_identifier(user.bank_account_number),
            'masked_bank_identifier': self.mask_identifier(user.bank_identifier),
            'masked_tax_id': self.mask_identifier(user.tax_id),
        }
        return render(request, self.template_name, context)

    @staticmethod
    def mask_identifier(value):
        if not value:
            return ''
        return f"•••• {value[-4:]}" if len(value) > 4 else '••••'

    def post(self, request):
        form = EmployeeProfileForm(request.POST, request.FILES, instance=request.user)
        if not form.is_valid():
            messages.error(request, 'Please correct the highlighted profile fields.')
            return self.render_profile(request, form)

        user = form.save(commit=False)
        user.profile_updated_at = timezone.now()
        user.save()
        
        messages.success(request, "Profile updated successfully!")
        return redirect('employee_profile')


class EmployeeProfilePhotoView(LoginRequiredMixin, View):
    def get(self, request, user_id=None):
        photo_user = request.user if user_id is None else get_object_or_404(User, pk=user_id)
        can_view_project_teammate = Project.objects.filter(
            Q(employees=request.user) | Q(assignments__employee=request.user)
        ).filter(
            Q(employees=photo_user) | Q(assignments__employee=photo_user)
        ).exists()
        if photo_user.pk != request.user.pk and request.user.role != 'admin' and not can_view_project_teammate:
            raise Http404
        if not photo_user.profile_photo:
            raise Http404
        content_type = mimetypes.guess_type(photo_user.profile_photo.name)[0] or 'application/octet-stream'
        return FileResponse(photo_user.profile_photo.open('rb'), content_type=content_type)


class EmployeeDocumentDownloadView(EmployeeRoleRequiredMixin, View):
    def get(self, request, pk):
        document = get_object_or_404(EmployeeDocument, pk=pk, employee=request.user)
        return FileResponse(
            document.file.open('rb'),
            as_attachment=True,
            filename=Path(document.file.name).name,
            content_type='application/octet-stream',
        )


class EmployeePasswordChangeView(EmployeeRoleRequiredMixin, PasswordChangeView):
    template_name = 'core/change_password.html'
    success_url = reverse_lazy('employee_profile')


class EmployeeSalaryView(EmployeeRoleRequiredMixin, TemplateView):
    """
    Displays current month's salary details and payment history.
    """
    template_name = 'core/salary.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        today = timezone.localdate()
        
        # Current month salary record
        current_salary = Salary.objects.filter(
            employee=user,
            month=today.month,
            year=today.year
        ).first()
        
        if current_salary:
            current_status = current_salary.status.title()
            current_amount = current_salary.amount
        else:
            current_status = "Pending"
            current_amount = user.salary_amount
            
        # Next salary due date
        next_salary_date = user.get_next_salary_due_date()
        
        # Salary history table
        salary_history = Salary.objects.filter(employee=user).order_by('-year', '-month')
        
        context.update({
            'current_status': current_status,
            'current_amount': current_amount,
            'next_salary_date': next_salary_date,
            'salary_history': salary_history,
        })
        return context


class EmployeeCalendarView(EmployeeRoleRequiredMixin, TemplateView):
    """
    Month-grid calendar combining the employee's own attendance status
    and leave applications (approved/pending) for the selected month.
    """
    template_name = 'core/calendar.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        today = timezone.localdate()

        month_str = self.request.GET.get('month')
        year_str = self.request.GET.get('year')
        selected_month = int(month_str) if month_str and month_str.isdigit() and 1 <= int(month_str) <= 12 else today.month
        selected_year = int(year_str) if year_str and year_str.isdigit() and 2000 <= int(year_str) <= today.year + 3 else today.year
        calendar_view = self.request.GET.get('view', 'month')
        if calendar_view not in {'month', 'week', 'agenda'}:
            calendar_view = 'month'

        month_start = datetime.date(selected_year, selected_month, 1)
        month_end = datetime.date(selected_year, selected_month, calendar.monthrange(selected_year, selected_month)[1])

        attendance_map = {
            att.date.day: att
            for att in Attendance.objects.filter(employee=user, date__year=selected_year, date__month=selected_month)
        }

        leave_map = {}
        leaves = LeaveApplication.objects.filter(
            employee=user,
            status__in=['pending', 'approved'],
            start_date__lte=month_end,
            end_date__gte=month_start,
        ).order_by('status')
        for leave in leaves:
            day = max(leave.start_date, month_start)
            end = min(leave.end_date, month_end)
            while day <= end:
                if day.day not in leave_map or leave.status == 'approved':
                    leave_map[day.day] = leave
                day += datetime.timedelta(days=1)

        holidays_map = {
            holiday.date.day: holiday
            for holiday in CompanyHoliday.objects.filter(date__range=(month_start, month_end))
        }
        team_leave_counts = {}
        team_leaves = LeaveApplication.objects.filter(
            employee__is_active=True,
            employee__role='employee',
            status='approved',
            start_date__lte=month_end,
            end_date__gte=month_start,
        ).exclude(employee=user).select_related('employee')
        for leave in team_leaves:
            day = max(leave.start_date, month_start)
            end = min(leave.end_date, month_end)
            while day <= end:
                if day.weekday() < 5:
                    team_leave_counts[day.day] = team_leave_counts.get(day.day, 0) + 1
                day += datetime.timedelta(days=1)

        birthday_map = {}
        teammates = User.objects.filter(role='employee', is_active=True, date_of_birth__isnull=False).exclude(pk=user.pk)
        for teammate in teammates:
            if teammate.date_of_birth.month == selected_month and teammate.date_of_birth.day <= calendar.monthrange(selected_year, selected_month)[1]:
                birthday_map.setdefault(teammate.date_of_birth.day, []).append(teammate.get_full_name() or teammate.username)

        cal = calendar.Calendar(firstweekday=6)  # Sunday-first
        weeks = []
        for week in cal.monthdayscalendar(selected_year, selected_month):
            week_data = []
            for day in week:
                if day == 0:
                    week_data.append(None)
                else:
                    day_date = datetime.date(selected_year, selected_month, day)
                    attendance = attendance_map.get(day)
                    leave = leave_map.get(day)
                    holiday = holidays_map.get(day)
                    net_seconds = get_attendance_work_seconds(attendance) if attendance else 0
                    week_data.append({
                        'day': day,
                        'date': day_date,
                        'date_iso': day_date.isoformat(),
                        'is_today': day_date == today,
                        'is_weekend': day_date.weekday() >= 5,
                        'attendance': attendance,
                        'attendance_status': attendance.status if attendance else '',
                        'check_in': attendance.check_in if attendance else None,
                        'check_out': attendance.check_out if attendance else None,
                        'working_time': format_duration(net_seconds) if net_seconds else '',
                        'leave': leave,
                        'leave_status': leave.status if leave else '',
                        'leave_type': leave.get_leave_type_display() if leave else '',
                        'leave_approver': (leave.approved_by.get_full_name() or leave.approved_by.username) if leave and leave.approved_by else '',
                        'holiday': holiday,
                        'team_away_count': team_leave_counts.get(day, 0),
                        'birthdays': birthday_map.get(day, []),
                    })
            weeks.append(week_data)

        if calendar_view == 'week':
            anchor = today if today.year == selected_year and today.month == selected_month else month_start
            weeks = [week for week in weeks if any(day and day['date'] == anchor for day in week)]
        elif calendar_view == 'agenda':
            agenda_days = [day for week in weeks for day in week if day and (day['attendance'] or day['leave'] or day['holiday'] or day['birthdays'])]
            context['agenda_days'] = agenda_days

        if selected_month == 1:
            prev_month, prev_year = 12, selected_year - 1
        else:
            prev_month, prev_year = selected_month - 1, selected_year
        if selected_month == 12:
            next_month, next_year = 1, selected_year + 1
        else:
            next_month, next_year = selected_month + 1, selected_year

        context.update({
            'weeks': weeks,
            'selected_month': selected_month,
            'selected_year': selected_year,
            'month_name': calendar.month_name[selected_month],
            'prev_month': prev_month,
            'prev_year': prev_year,
            'next_month': next_month,
            'next_year': next_year,
            'today': today,
            'calendar_view': calendar_view,
            'leave_types': [choice[0] for choice in LeaveApplication.LEAVE_TYPE_CHOICES],
        })
        return context


class EmployeeProjectsView(EmployeeRoleRequiredMixin, TemplateView):
    """
    Lists the projects the logged-in employee is currently working on.
    """
    template_name = 'core/projects.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        selected_filter = self.request.GET.get('status', 'all').lower()
        if selected_filter not in {'all', 'active', 'completed', 'on_hold'}:
            selected_filter = 'all'

        employee_projects = Project.objects.filter(
            Q(employees=user) | Q(assignments__employee=user)
        ).distinct()
        projects = employee_projects.prefetch_related(
            'employees',
            'assignments__employee__manager',
            'milestones',
            'tasks',
            'technologies',
            'resources',
        ).order_by('name')
        if selected_filter != 'all':
            projects = projects.filter(status=selected_filter)

        project_details = []
        for project in projects:
            assignments = list(project.assignments.all())
            user_assignment = next((assignment for assignment in assignments if assignment.employee_id == user.pk), None)
            members = {member.pk: member for member in project.employees.all()}
            members.update({assignment.employee_id: assignment.employee for assignment in assignments})
            milestones = list(project.milestones.all())
            completed_milestones = sum(milestone.status == 'done' for milestone in milestones)
            progress = round(completed_milestones * 100 / len(milestones)) if milestones else 0
            project.team_members = sorted(
                members.values(),
                key=lambda item: (item.first_name, item.last_name, item.username),
            )
            project.employee_role = user_assignment.role if user_assignment else 'Contributor'
            project.progress = progress
            team = []
            for member in project.team_members:
                full_name = member.get_full_name() or member.username
                initials = ''.join(part[0] for part in full_name.split()[:2]).upper()
                team.append({
                    'id': member.pk,
                    'name': full_name,
                    'designation': member.designation or 'Team member',
                    'manager': member.manager.get_full_name() if member.manager else '',
                    'avatar_url': reverse('employee_photo_by_id', args=[member.pk]) if member.profile_photo else '',
                    'initials': initials or member.username[:2].upper(),
                })

            project_details.append({
                'id': project.pk,
                'name': project.name,
                'category': project.category,
                'description': project.description,
                'objective': project.objective,
                'goals': project.goals,
                'scope': project.scope,
                'status': project.status,
                'status_label': project.get_status_display(),
                'role': user_assignment.role if user_assignment else 'Contributor',
                'start_date': project.start_date.strftime('%b %d, %Y') if project.start_date else '-',
                'end_date': project.end_date.strftime('%b %d, %Y') if project.end_date else '-',
                'progress': progress,
                'milestones': [{
                    'title': milestone.title,
                    'description': milestone.description,
                    'due_date': milestone.due_date.strftime('%b %d, %Y') if milestone.due_date else '',
                    'status': milestone.status,
                    'status_label': milestone.get_status_display(),
                } for milestone in milestones],
                'tasks': [{
                    'title': task.title,
                    'status': task.status,
                    'status_label': task.get_status_display(),
                    'due_date': task.due_date.strftime('%b %d, %Y') if task.due_date else '',
                } for task in project.tasks.all() if task.employee_id == user.pk],
                'technologies': [technology.name for technology in project.technologies.all()],
                'resources': [{
                    'title': resource.title,
                    'type': resource.get_resource_type_display(),
                    'url': resource.url,
                } for resource in project.resources.all()],
                'team': team,
            })

        context.update({
            'projects': projects,
            'project_details': project_details,
            'selected_filter': selected_filter,
            'project_filters': [
                {'value': 'all', 'label': 'All Projects', 'count': employee_projects.count()},
                {'value': 'active', 'label': 'Active Projects', 'count': employee_projects.filter(status='active').count()},
                {'value': 'completed', 'label': 'Completed Projects', 'count': employee_projects.filter(status='completed').count()},
                {'value': 'on_hold', 'label': 'On Hold', 'count': employee_projects.filter(status='on_hold').count()},
            ],
        })
        return context

